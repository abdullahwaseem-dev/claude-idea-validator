#!/usr/bin/env python3
"""
Generate a PDF and a Word (.docx) market-research report from a JSON file
matching the schema documented in references/report_schema.md.

Usage:
    python generate_report.py --input report.json --outdir ./output --name my-idea

Produces:
    ./output/my-idea.pdf
    ./output/my-idea.docx
"""

import argparse
import json
import sys
from pathlib import Path

VERDICT_COLORS = {
    "VALIDATED": "2E7D32",   # green
    "CONDITIONAL": "E65100",  # amber
    "PIVOT": "1565C0",       # blue
}


def load_report(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def score_label(dimension_key: str) -> str:
    labels = {
        "market_size": "Market Size & Growth",
        "competition_intensity": "Competition Intensity",
        "differentiation": "Differentiation / Moat",
        "feasibility": "Execution Feasibility",
        "timing": "Timing",
        "monetization": "Monetization Potential",
    }
    return labels.get(dimension_key, dimension_key.replace("_", " ").title())


# --------------------------------------------------------------------------
# PDF generation (reportlab)
# --------------------------------------------------------------------------

def build_pdf(report: dict, out_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable,
        ListItem, HRFlowable,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H1Custom", parent=styles["Heading1"], spaceBefore=18, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2Custom", parent=styles["Heading2"], spaceBefore=14, spaceAfter=6))
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], spaceAfter=8, leading=15))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=9, textColor=colors.HexColor("#555555")))

    meta = report.get("meta", {})
    verdict = meta.get("verdict", "").upper()
    verdict_color = colors.HexColor("#" + VERDICT_COLORS.get(verdict, "333333"))
    styles.add(ParagraphStyle(name="Verdict", parent=styles["Heading2"], textColor=verdict_color))

    story = []

    story.append(Paragraph(meta.get("idea_title", "Market Research Report"), styles["Title"]))
    story.append(Paragraph(meta.get("one_line_pitch", ""), styles["BodyCustom"]))
    story.append(Paragraph(meta.get("date", ""), styles["Small"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"Verdict: {verdict}", styles["Verdict"]))
    story.append(Paragraph(meta.get("verdict_headline", ""), styles["BodyCustom"]))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))

    def bullets(items):
        return ListFlowable(
            [ListItem(Paragraph(str(i), styles["BodyCustom"])) for i in items],
            bulletType="bullet",
        )

    # Executive summary
    story.append(Paragraph("Executive Summary", styles["H1Custom"]))
    for para in report.get("executive_summary", "").split("\n\n"):
        if para.strip():
            story.append(Paragraph(para.strip(), styles["BodyCustom"]))

    # Market landscape
    ml = report.get("market_landscape", {})
    if ml:
        story.append(Paragraph("Market Landscape", styles["H1Custom"]))
        if ml.get("overview"):
            story.append(Paragraph(ml["overview"], styles["BodyCustom"]))
        if ml.get("market_size_estimate"):
            story.append(Paragraph(f"<b>Market size estimate:</b> {ml['market_size_estimate']}", styles["BodyCustom"]))
        if ml.get("growth_trends"):
            story.append(Paragraph("Growth trends:", styles["BodyCustom"]))
            story.append(bullets(ml["growth_trends"]))
        if ml.get("demand_signals"):
            story.append(Paragraph("Demand signals:", styles["BodyCustom"]))
            story.append(bullets(ml["demand_signals"]))
        conf = ml.get("data_confidence")
        if conf:
            story.append(Paragraph(f"<b>Data confidence:</b> {conf.upper()}", styles["Small"]))
        if ml.get("data_limitations"):
            story.append(Paragraph(f"<i>Data limitations:</i> {ml['data_limitations']}", styles["Small"]))

    # Competitors
    competitors = report.get("competitors", [])
    if competitors:
        story.append(Paragraph("Competitor Analysis", styles["H1Custom"]))
        for c in competitors:
            story.append(Paragraph(f"{c.get('name','')} <font color='#777777'>({c.get('type','')})</font>", styles["H2Custom"]))
            if c.get("description"):
                story.append(Paragraph(c["description"], styles["BodyCustom"]))
            if c.get("strengths"):
                story.append(Paragraph("Strengths:", styles["BodyCustom"]))
                story.append(bullets(c["strengths"]))
            if c.get("weaknesses"):
                story.append(Paragraph("Weaknesses:", styles["BodyCustom"]))
                story.append(bullets(c["weaknesses"]))
            if c.get("pricing_positioning"):
                story.append(Paragraph(f"<b>Pricing / positioning:</b> {c['pricing_positioning']}", styles["BodyCustom"]))

    # Target audience
    ta = report.get("target_audience", {})
    if ta:
        story.append(Paragraph("Target Audience", styles["H1Custom"]))
        if ta.get("primary_segment"):
            story.append(Paragraph(f"<b>Primary segment:</b> {ta['primary_segment']}", styles["BodyCustom"]))
        if ta.get("secondary_segments"):
            story.append(Paragraph("Secondary segments:", styles["BodyCustom"]))
            story.append(bullets(ta["secondary_segments"]))
        if ta.get("pain_points"):
            story.append(Paragraph("Pain points:", styles["BodyCustom"]))
            story.append(bullets(ta["pain_points"]))
        if ta.get("willingness_to_pay"):
            story.append(Paragraph(f"<b>Willingness to pay:</b> {ta['willingness_to_pay']}", styles["BodyCustom"]))
        if ta.get("acquisition_channels"):
            story.append(Paragraph("Acquisition channels:", styles["BodyCustom"]))
            story.append(bullets(ta["acquisition_channels"]))

    # Viability scorecard
    vs = report.get("viability_scores", {})
    if vs:
        story.append(Paragraph("Viability Scorecard", styles["H1Custom"]))
        dim_keys = ["market_size", "competition_intensity", "differentiation", "feasibility", "timing", "monetization"]
        table_data = [["Dimension", "Score /5", "Rationale"]]
        for k in dim_keys:
            entry = vs.get(k)
            if entry:
                table_data.append([score_label(k), str(entry.get("score", "")), Paragraph(entry.get("rationale", ""), styles["BodyCustom"])])
        t = Table(table_data, colWidths=[1.6 * inch, 0.7 * inch, 3.7 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t)
        story.append(Spacer(1, 8))
        if "composite" in vs:
            story.append(Paragraph(
                f"<b>Composite score: {vs['composite']} / 5 &nbsp;→&nbsp; {vs.get('composite_band','')}</b>",
                styles["BodyCustom"],
            ))

    # Opportunities & Risks
    opp = report.get("opportunities", [])
    risks = report.get("risks", [])
    if opp or risks:
        story.append(Paragraph("Opportunity & Risk Assessment", styles["H1Custom"]))
        if opp:
            story.append(Paragraph("Opportunities:", styles["H2Custom"]))
            story.append(bullets(opp))
        if risks:
            story.append(Paragraph("Risks:", styles["H2Custom"]))
            story.append(bullets(risks))

    # Recommendation
    rec = report.get("recommendation", {})
    if rec:
        story.append(Paragraph("Recommendation", styles["H1Custom"]))
        story.append(Paragraph(f"<b>Verdict: {rec.get('verdict','')}</b>", styles["Verdict"]))
        if rec.get("rationale"):
            story.append(Paragraph(rec["rationale"], styles["BodyCustom"]))
        if rec.get("suggested_refinements"):
            story.append(Paragraph("Suggested refinements:", styles["BodyCustom"]))
            story.append(bullets(rec["suggested_refinements"]))
        alternatives = rec.get("alternatives", [])
        if alternatives:
            story.append(Paragraph("Alternative Directions", styles["H2Custom"]))
            for alt in alternatives:
                story.append(Paragraph(f"<b>{alt.get('name','')}</b> — {alt.get('pitch','')}", styles["BodyCustom"]))
                if alt.get("why_better"):
                    story.append(Paragraph(f"<i>Why this is a better bet:</i> {alt['why_better']}", styles["BodyCustom"]))
                if alt.get("evidence"):
                    story.append(Paragraph(f"<i>Evidence:</i> {alt['evidence']}", styles["BodyCustom"]))

    # Sources
    sources = report.get("sources", [])
    if sources:
        story.append(Paragraph("Sources", styles["H1Custom"]))
        for s in sources:
            label = s.get("label", "")
            url = s.get("url", "")
            story.append(Paragraph(f'{label} — <link href="{url}" color="blue">{url}</link>', styles["Small"]))

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=LETTER,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        title=meta.get("idea_title", "Market Research Report"),
    )
    doc.build(story)


# --------------------------------------------------------------------------
# DOCX generation (python-docx)
# --------------------------------------------------------------------------

def build_docx(report: dict, out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    meta = report.get("meta", {})
    verdict = meta.get("verdict", "").upper()
    hexcolor = VERDICT_COLORS.get(verdict, "333333")
    rgb = RGBColor(int(hexcolor[0:2], 16), int(hexcolor[2:4], 16), int(hexcolor[4:6], 16))

    doc = Document()

    doc.add_heading(meta.get("idea_title", "Market Research Report"), level=0)
    p = doc.add_paragraph(meta.get("one_line_pitch", ""))
    p.runs[0].italic = True
    date_p = doc.add_paragraph(meta.get("date", ""))
    date_p.runs[0].font.size = Pt(9)
    date_p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    verdict_p = doc.add_paragraph()
    run = verdict_p.add_run(f"Verdict: {verdict}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = rgb
    doc.add_paragraph(meta.get("verdict_headline", ""))

    def add_bullets(items):
        for item in items:
            doc.add_paragraph(str(item), style="List Bullet")

    doc.add_heading("Executive Summary", level=1)
    for para in report.get("executive_summary", "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    ml = report.get("market_landscape", {})
    if ml:
        doc.add_heading("Market Landscape", level=1)
        if ml.get("overview"):
            doc.add_paragraph(ml["overview"])
        if ml.get("market_size_estimate"):
            p = doc.add_paragraph()
            p.add_run("Market size estimate: ").bold = True
            p.add_run(ml["market_size_estimate"])
        if ml.get("growth_trends"):
            doc.add_paragraph("Growth trends:")
            add_bullets(ml["growth_trends"])
        if ml.get("demand_signals"):
            doc.add_paragraph("Demand signals:")
            add_bullets(ml["demand_signals"])
        if ml.get("data_confidence"):
            p = doc.add_paragraph()
            p.add_run(f"Data confidence: {ml['data_confidence'].upper()}").italic = True
        if ml.get("data_limitations"):
            p = doc.add_paragraph()
            p.add_run(f"Data limitations: {ml['data_limitations']}").italic = True

    competitors = report.get("competitors", [])
    if competitors:
        doc.add_heading("Competitor Analysis", level=1)
        for c in competitors:
            doc.add_heading(f"{c.get('name','')} ({c.get('type','')})", level=2)
            if c.get("description"):
                doc.add_paragraph(c["description"])
            if c.get("strengths"):
                doc.add_paragraph("Strengths:")
                add_bullets(c["strengths"])
            if c.get("weaknesses"):
                doc.add_paragraph("Weaknesses:")
                add_bullets(c["weaknesses"])
            if c.get("pricing_positioning"):
                p = doc.add_paragraph()
                p.add_run("Pricing / positioning: ").bold = True
                p.add_run(c["pricing_positioning"])

    ta = report.get("target_audience", {})
    if ta:
        doc.add_heading("Target Audience", level=1)
        if ta.get("primary_segment"):
            p = doc.add_paragraph()
            p.add_run("Primary segment: ").bold = True
            p.add_run(ta["primary_segment"])
        if ta.get("secondary_segments"):
            doc.add_paragraph("Secondary segments:")
            add_bullets(ta["secondary_segments"])
        if ta.get("pain_points"):
            doc.add_paragraph("Pain points:")
            add_bullets(ta["pain_points"])
        if ta.get("willingness_to_pay"):
            p = doc.add_paragraph()
            p.add_run("Willingness to pay: ").bold = True
            p.add_run(ta["willingness_to_pay"])
        if ta.get("acquisition_channels"):
            doc.add_paragraph("Acquisition channels:")
            add_bullets(ta["acquisition_channels"])

    vs = report.get("viability_scores", {})
    if vs:
        doc.add_heading("Viability Scorecard", level=1)
        dim_keys = ["market_size", "competition_intensity", "differentiation", "feasibility", "timing", "monetization"]
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Dimension", "Score /5", "Rationale"
        for k in dim_keys:
            entry = vs.get(k)
            if entry:
                row = table.add_row().cells
                row[0].text = score_label(k)
                row[1].text = str(entry.get("score", ""))
                row[2].text = entry.get("rationale", "")
        if "composite" in vs:
            p = doc.add_paragraph()
            run = p.add_run(f"Composite score: {vs['composite']} / 5  →  {vs.get('composite_band','')}")
            run.bold = True

    opp = report.get("opportunities", [])
    risks = report.get("risks", [])
    if opp or risks:
        doc.add_heading("Opportunity & Risk Assessment", level=1)
        if opp:
            doc.add_heading("Opportunities", level=2)
            add_bullets(opp)
        if risks:
            doc.add_heading("Risks", level=2)
            add_bullets(risks)

    rec = report.get("recommendation", {})
    if rec:
        doc.add_heading("Recommendation", level=1)
        p = doc.add_paragraph()
        run = p.add_run(f"Verdict: {rec.get('verdict','')}")
        run.bold = True
        run.font.color.rgb = rgb
        if rec.get("rationale"):
            doc.add_paragraph(rec["rationale"])
        if rec.get("suggested_refinements"):
            doc.add_paragraph("Suggested refinements:")
            add_bullets(rec["suggested_refinements"])
        alternatives = rec.get("alternatives", [])
        if alternatives:
            doc.add_heading("Alternative Directions", level=2)
            for alt in alternatives:
                p = doc.add_paragraph()
                p.add_run(f"{alt.get('name','')} — ").bold = True
                p.add_run(alt.get("pitch", ""))
                if alt.get("why_better"):
                    p2 = doc.add_paragraph()
                    p2.add_run("Why this is a better bet: ").italic = True
                    p2.add_run(alt["why_better"])
                if alt.get("evidence"):
                    p3 = doc.add_paragraph()
                    p3.add_run("Evidence: ").italic = True
                    p3.add_run(alt["evidence"])

    sources = report.get("sources", [])
    if sources:
        doc.add_heading("Sources", level=1)
        for s in sources:
            doc.add_paragraph(f"{s.get('label','')} — {s.get('url','')}")

    doc.save(str(out_path))


def main():
    parser = argparse.ArgumentParser(description="Generate a PDF + DOCX market research report from a JSON spec.")
    parser.add_argument("--input", required=True, help="Path to the report JSON file")
    parser.add_argument("--outdir", default="./output", help="Directory to write the PDF/DOCX into")
    parser.add_argument("--name", required=True, help="Filesystem-safe base name for the output files (no extension)")
    args = parser.parse_args()

    report = load_report(args.input)
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    pdf_path = outdir / f"{args.name}.pdf"
    docx_path = outdir / f"{args.name}.docx"

    try:
        build_pdf(report, pdf_path)
        print(f"Wrote {pdf_path}")
    except ImportError:
        print("reportlab is not installed — run: pip install -r requirements.txt", file=sys.stderr)
        raise

    try:
        build_docx(report, docx_path)
        print(f"Wrote {docx_path}")
    except ImportError:
        print("python-docx is not installed — run: pip install -r requirements.txt", file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
